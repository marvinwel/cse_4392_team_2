import boto3
import csv
import hashlib
import io
import json
import logging
import os
import time
import urllib.error
import urllib.request
import uuid

from datetime import datetime, timezone
from pathlib import PurePosixPath
from urllib.parse import unquote_plus

from botocore.exceptions import ClientError
from pypdf import PdfReader
from docx import Document

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_aws import BedrockEmbeddings


# ==========================================================
# CONFIGURATION
# ==========================================================

logger = logging.getLogger()
logger.setLevel(logging.INFO)

AWS_REGION = "us-east-2"

EMBEDDING_MODEL_ID = "amazon.titan-embed-text-v2:0"

CHROMA_URL = os.environ.get(
    "CHROMA_URL",
    "http://18.224.228.207:8000"
).rstrip("/")

CHROMA_COLLECTION_ID = os.environ.get(
    "CHROMA_COLLECTION_ID",
    "2700ece6-2519-4fc4-8db0-f72353e5ac63"
)

DYNAMODB_TABLE = "uta-ai-advisor-documents"

TEAMS_WEBHOOK_URL = os.environ.get("TEAMS_WEBHOOK_URL")

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".pdf", ".docx"
}

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150

LOCK_SECONDS = 900


# ==========================================================
# AWS CLIENTS
# ==========================================================

s3 = boto3.client("s3", region_name=AWS_REGION)

bedrock = boto3.client(
    "bedrock-runtime",
    region_name=AWS_REGION
)

dynamodb = boto3.resource(
    "dynamodb",
    region_name=AWS_REGION
)

documents_table = dynamodb.Table(DYNAMODB_TABLE)


# ==========================================================
# LANGCHAIN COMPONENTS
# ==========================================================

# Splits at paragraph, then line, then sentence, then word boundaries
# before falling back to a hard cut, so chunks rarely end mid-sentence.
splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
    separators=["\n\n", "\n", ". ", " ", ""],
)

# Same model, dimensions, and normalization as before, so new vectors
# stay compatible with the ones already stored in ChromaDB.
embeddings = BedrockEmbeddings(
    model_id=EMBEDDING_MODEL_ID,
    region_name=AWS_REGION,
    client=bedrock,
    model_kwargs={
        "dimensions": 1024,
        "normalize": True
    },
)


# ==========================================================
# COMMON HELPERS
# ==========================================================

def utc_now():
    return datetime.now(timezone.utc).isoformat()


def document_id_for(bucket, key):
    return f"{bucket}/{key}"


def normalize_etag(etag):
    return (etag or "").strip('"')


def chroma_endpoint(action):
    return (
        f"{CHROMA_URL}/api/v2/tenants/default_tenant/"
        f"databases/default_database/collections/"
        f"{CHROMA_COLLECTION_ID}/{action}"
    )


def chroma_request(action, payload):
    request = urllib.request.Request(
        chroma_endpoint(action),
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST"
    )

    with urllib.request.urlopen(request, timeout=30) as response:
        body = response.read()

        if response.status not in (200, 201):
            raise RuntimeError(
                f"ChromaDB HTTP {response.status}"
            )

        return json.loads(body) if body else {}


# ==========================================================
# DOCUMENT EXTRACTION
# ==========================================================

def extract_text(filename, file_bytes):
    extension = PurePosixPath(filename).suffix.lower()

    if extension in (".txt", ".md"):
        return file_bytes.decode("utf-8-sig")

    if extension == ".csv":
        content = file_bytes.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))

        return "\n".join(
            " | ".join(row)
            for row in reader
        )

    if extension == ".json":
        content = file_bytes.decode("utf-8-sig")
        data = json.loads(content)

        return json.dumps(
            data,
            indent=2,
            ensure_ascii=False
        )

    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(file_bytes))

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if extension == ".docx":
        document = Document(io.BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        tables = []

        for table in document.tables:
            for row in table.rows:
                tables.append(
                    " | ".join(
                        cell.text for cell in row.cells
                    )
                )

        return "\n".join(paragraphs + tables)

    raise ValueError(
        f"Unsupported document format: {extension}"
    )


# ==========================================================
# DYNAMODB DOCUMENT TRACKING
# ==========================================================

def get_document_record(document_id):
    response = documents_table.get_item(
        Key={"document_id": document_id},
        ConsistentRead=True
    )

    return response.get("Item")


def acquire_processing_lock(document_id, document_hash):
    now = int(time.time())
    lock_token = str(uuid.uuid4())

    try:
        documents_table.update_item(
            Key={"document_id": document_id},
            UpdateExpression=(
                "SET #status = :processing, "
                "lock_token = :token, "
                "lock_until = :lock_until, "
                "updated_at = :updated_at"
            ),
            ConditionExpression=(
                "(attribute_not_exists(document_id) "
                "OR #status <> :completed "
                "OR document_hash <> :hash) "
                "AND "
                "(attribute_not_exists(lock_until) "
                "OR lock_until < :now)"
            ),
            ExpressionAttributeNames={
                "#status": "status"
            },
            ExpressionAttributeValues={
                ":processing": "PROCESSING",
                ":completed": "COMPLETED",
                ":token": lock_token,
                ":lock_until": now + LOCK_SECONDS,
                ":updated_at": utc_now(),
                ":hash": document_hash,
                ":now": now
            }
        )

        return lock_token

    except ClientError as error:
        if error.response["Error"]["Code"] == (
            "ConditionalCheckFailedException"
        ):
            return None

        raise


def mark_document_completed(
    document_id,
    document_hash,
    bucket,
    key,
    etag,
    chunk_count,
    lock_token
):
    documents_table.update_item(
        Key={"document_id": document_id},
        UpdateExpression=(
            "SET #status = :completed, "
            "document_hash = :hash, "
            "s3_etag = :etag, "
            "s3_bucket = :bucket, "
            "s3_key = :key, "
            "chunk_count = :chunk_count, "
            "last_processed = :processed "
            "REMOVE lock_token, lock_until"
        ),
        ConditionExpression=(
            "lock_token = :token "
            "AND #status = :processing"
        ),
        ExpressionAttributeNames={
            "#status": "status"
        },
        ExpressionAttributeValues={
            ":completed": "COMPLETED",
            ":processing": "PROCESSING",
            ":hash": document_hash,
            ":etag": etag,
            ":bucket": bucket,
            ":key": key,
            ":chunk_count": chunk_count,
            ":processed": utc_now(),
            ":token": lock_token
        }
    )


def update_unchanged_document_etag(document_id, etag):
    documents_table.update_item(
        Key={"document_id": document_id},
        UpdateExpression=(
            "SET s3_etag = :etag, "
            "last_checked = :checked"
        ),
        ConditionExpression="#status = :completed",
        ExpressionAttributeNames={
            "#status": "status"
        },
        ExpressionAttributeValues={
            ":etag": etag,
            ":checked": utc_now(),
            ":completed": "COMPLETED"
        }
    )


def mark_document_failed(document_id, lock_token, error):
    try:
        documents_table.update_item(
            Key={"document_id": document_id},
            UpdateExpression=(
                "SET #status = :failed, "
                "error_message = :error, "
                "updated_at = :updated_at "
                "REMOVE lock_token, lock_until"
            ),
            ConditionExpression="lock_token = :token",
            ExpressionAttributeNames={
                "#status": "status"
            },
            ExpressionAttributeValues={
                ":failed": "FAILED",
                ":error": str(error)[:1000],
                ":updated_at": utc_now(),
                ":token": lock_token
            }
        )
    except Exception:
        logger.exception(
            "Unable to mark document FAILED: %s",
            document_id
        )


# ==========================================================
# CHROMADB OPERATIONS
# ==========================================================

def get_existing_document_chunks(source):
    result = chroma_request(
        "get",
        {
            "where": {"source": source},
            "include": ["metadatas"]
        }
    )

    return result.get("ids", [])


def save_to_chromadb(
    chunk_id,
    text,
    embedding,
    source,
    chunk_index,
    document_hash
):
    chroma_request(
        "upsert",
        {
            "ids": [chunk_id],
            "embeddings": [embedding],
            "documents": [text],
            "metadatas": [{
                "source": source,
                "chunk_index": chunk_index,
                "document_hash": document_hash
            }]
        }
    )

    logger.info(
        "Stored chunk %s in ChromaDB",
        chunk_id
    )


def delete_chromadb_chunks(chunk_ids):
    if not chunk_ids:
        return

    batch_size = 100

    for start in range(0, len(chunk_ids), batch_size):
        batch = chunk_ids[start:start + batch_size]

        chroma_request(
            "delete",
            {"ids": batch}
        )

    logger.info(
        "Deleted %d outdated ChromaDB chunks",
        len(chunk_ids)
    )


# ==========================================================
# MICROSOFT TEAMS NOTIFICATION
# ==========================================================

def notify_teams(filename, chunk_count):
    if not TEAMS_WEBHOOK_URL:
        logger.warning(
            "TEAMS_WEBHOOK_URL is not configured"
        )
        return

    message = {
        "type": "message",
        "attachments": [
            {
                "contentType":
                    "application/vnd.microsoft.card.adaptive",
                "content": {
                    "$schema":
                        "http://adaptivecards.io/schemas/adaptive-card.json",
                    "type": "AdaptiveCard",
                    "version": "1.2",
                    "body": [
                        {
                            "type": "TextBlock",
                            "text":
                                "UTA AI Advisor - Ingestion Successful",
                            "weight": "Bolder",
                            "size": "Medium"
                        },
                        {
                            "type": "FactSet",
                            "facts": [
                                {
                                    "title": "Document",
                                    "value": filename
                                },
                                {
                                    "title": "Chunks stored",
                                    "value": str(chunk_count)
                                },
                                {
                                    "title": "Embedding model",
                                    "value": "Amazon Titan V2"
                                },
                                {
                                    "title": "Vector database",
                                    "value": "ChromaDB"
                                },
                                {
                                    "title": "Status",
                                    "value": "Completed"
                                }
                            ]
                        }
                    ]
                }
            }
        ]
    }

    request = urllib.request.Request(
        TEAMS_WEBHOOK_URL,
        data=json.dumps(message).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST"
    )

    try:
        with urllib.request.urlopen(
            request, timeout=10
        ) as response:
            logger.info(
                "Teams webhook accepted: HTTP %s",
                response.status
            )

    except urllib.error.HTTPError as error:
        logger.error(
            "Teams notification failed: HTTP %s",
            error.code
        )

    except Exception:
        logger.exception(
            "Teams notification failed"
        )


# ==========================================================
# MAIN LAMBDA HANDLER
# ==========================================================

def lambda_handler(event, context):
    processed = 0
    skipped = 0
    stored_chunks = 0

    for record in event.get("Records", []):
        if record.get("eventSource") != "aws:s3":
            continue

        bucket = record["s3"]["bucket"]["name"]
        key = unquote_plus(
            record["s3"]["object"]["key"]
        )

        extension = PurePosixPath(key).suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            logger.info(
                "Skipping unsupported file: %s",
                key
            )
            skipped += 1
            continue

        document_id = document_id_for(bucket, key)

        # ==================================================
        # STEP 1: CHECK DYNAMODB BY DOCUMENT NAME FIRST
        # ==================================================

        existing = get_document_record(document_id)

        event_etag = normalize_etag(
            record["s3"]["object"].get("eTag")
        )

        if (
            existing
            and existing.get("status") == "COMPLETED"
            and event_etag
            and existing.get("s3_etag") == event_etag
        ):
            logger.info(
                "SKIPPED: Document already processed "
                "with matching ETag: %s",
                key
            )
            skipped += 1
            continue

        logger.info(
            "Document requires further checking: %s",
            key
        )

        # ==================================================
        # STEP 2: DOWNLOAD ONLY IF NEEDED
        # ==================================================

        response = s3.get_object(
            Bucket=bucket,
            Key=key
        )

        file_bytes = response["Body"].read()

        actual_etag = normalize_etag(
            response.get("ETag")
        )

        document_hash = hashlib.sha256(
            file_bytes
        ).hexdigest()

        # ==================================================
        # STEP 3: CHECK WHETHER CONTENT CHANGED
        # ==================================================

        if (
            existing
            and existing.get("status") == "COMPLETED"
            and existing.get("document_hash") == document_hash
        ):
            logger.info(
                "SKIPPED: Identical document content: %s",
                key
            )

            update_unchanged_document_etag(
                document_id,
                actual_etag
            )

            skipped += 1
            continue

        # ==================================================
        # STEP 4: ACQUIRE PROCESSING LOCK
        # ==================================================

        lock_token = acquire_processing_lock(
            document_id,
            document_hash
        )

        if not lock_token:
            logger.info(
                "SKIPPED: Document already completed "
                "or currently processing: %s",
                key
            )
            skipped += 1
            continue

        try:
            # STEP 5: Extract text.
            text = extract_text(
                key,
                file_bytes
            )

            if not text.strip():
                raise ValueError(
                    f"No text extracted from {key}"
                )

            logger.info(
                "Extracted %d characters from %s",
                len(text),
                key
            )

            # STEP 6: Split text into chunks (LangChain).
            chunks = splitter.split_text(text)

            logger.info(
                "Created %d chunks from %s",
                len(chunks),
                key
            )

            # STEP 7: Get existing vector IDs.
            existing_ids = get_existing_document_chunks(
                key
            )

            # STEP 8: Generate embeddings only now (LangChain).
            vectors = embeddings.embed_documents(chunks)

            new_ids = []

            for index, (chunk, embedding) in enumerate(
                zip(chunks, vectors)
            ):
                chunk_id = hashlib.sha256(
                    (
                        f"{bucket}/{key}/"
                        f"{document_hash}/{index}"
                    ).encode("utf-8")
                ).hexdigest()

                save_to_chromadb(
                    chunk_id=chunk_id,
                    text=chunk,
                    embedding=embedding,
                    source=key,
                    chunk_index=index,
                    document_hash=document_hash
                )

                new_ids.append(chunk_id)
                stored_chunks += 1

                logger.info(
                    "Chunk %d stored: %d characters, "
                    "%d dimensions",
                    index + 1,
                    len(chunk),
                    len(embedding)
                )

            # STEP 9: Remove old vectors.
            new_id_set = set(new_ids)

            outdated_ids = [
                chunk_id
                for chunk_id in existing_ids
                if chunk_id not in new_id_set
            ]

            delete_chromadb_chunks(
                outdated_ids
            )

            # STEP 10: Update DynamoDB.
            mark_document_completed(
                document_id=document_id,
                document_hash=document_hash,
                bucket=bucket,
                key=key,
                etag=actual_etag,
                chunk_count=len(chunks),
                lock_token=lock_token
            )

            processed += 1

            logger.info(
                "SUCCESS: Document ingestion completed: %s",
                key
            )

        except Exception as error:
            logger.exception(
                "FAILED: Document ingestion: %s",
                key
            )

            mark_document_failed(
                document_id,
                lock_token,
                error
            )

            raise

        # STEP 11: Notify Teams.
        notify_teams(
            filename=key,
            chunk_count=len(chunks)
        )

    return {
        "statusCode": 200,
        "body": json.dumps({
            "processed_documents": processed,
            "skipped_documents": skipped,
            "stored_chunks": stored_chunks
        })
    }